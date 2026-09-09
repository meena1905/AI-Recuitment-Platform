import os
import tempfile

os.environ["DATABASE_URL"] = "sqlite:///./test_talenta.db"
os.environ.setdefault("JWT_SECRET_KEY", "test-secret-key-for-pytest")

import pytest
from docx import Document
from fastapi.testclient import TestClient

from main import app, get_db
from resume_parser import extract_text_from_file

client = TestClient(app)

HR = {"name": "Test HR", "email": "hr@test.com", "password": "secret123", "role": "hr", "company_name": "TestCorp"}
CANDIDATE = {"name": "Alex Doe", "email": "alex@test.com", "password": "secret123", "role": "candidate"}


@pytest.fixture(autouse=True)
def clean_db():
    from database import engine
    from sqlalchemy import text
    with engine.begin() as connection:
        for table in ({"applications", "interviews", "jobs", "users", "companies"} & {
            r[0] for r in connection.execute(text("SELECT name FROM sqlite_master WHERE type='table'"))
        }):
            connection.execute(text(f"DELETE FROM {table}"))
    yield


def register(user_data):
    return client.post("/auth/register", json=user_data)


def login(email, password="secret123"):
    return client.post("/auth/login", json={"email": email, "password": password})


@pytest.fixture()
def hr_token():
    register(HR)
    return login(HR["email"]).json()["access_token"]


@pytest.fixture()
def candidate_token():
    register(CANDIDATE)
    return login(CANDIDATE["email"]).json()["access_token"]


@pytest.fixture()
def job_id(hr_token):
    response = client.post(
        "/jobs",
        json={"title": "Backend Engineer", "description": "Build APIs", "requirements": "Python, FastAPI"},
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    return response.json()["id"]


@pytest.fixture()
def sample_resume(tmp_path):
    doc = Document()
    doc.add_paragraph("Alex Doe")
    doc.add_paragraph("Python developer with experience in FastAPI, SQL, and REST APIs.")
    path = tmp_path / "resume.docx"
    doc.save(path)
    return path


# --- Tests ---

def test_user_registration():
    response = register({"name": "New User", "email": "new@test.com", "password": "pw123456", "role": "candidate"})
    assert response.status_code == 200
    assert response.json()["email"] == "new@test.com"


def test_login():
    register(HR)
    response = login(HR["email"])
    assert response.status_code == 200
    assert "access_token" in response.json()


def test_login_invalid_password():
    register(HR)
    response = login(HR["email"], password="wrongpassword")
    assert response.status_code == 401
    assert "Invalid email or password" in response.json()["detail"]


def test_create_job(hr_token):
    response = client.post(
        "/jobs",
        json={"title": "Data Scientist", "description": "ML", "requirements": "Python, ML"},
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    assert response.json()["title"] == "Data Scientist"


def test_apply_to_job(job_id, candidate_token, sample_resume):
    with open(sample_resume, "rb") as resume:
        response = client.post(
            f"/jobs/{job_id}/apply",
            files={"resume": ("resume.docx", resume, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers={"Authorization": f"Bearer {candidate_token}"},
        )
    assert response.status_code == 200
    assert response.json()["status"] == "applied"


def test_duplicate_application(job_id, candidate_token, sample_resume):
    with open(sample_resume, "rb") as resume:
        client.post(
            f"/jobs/{job_id}/apply",
            files={"resume": ("resume.docx", resume, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers={"Authorization": f"Bearer {candidate_token}"},
        )
    with open(sample_resume, "rb") as resume:
        response = client.post(
            f"/jobs/{job_id}/apply",
            files={"resume": ("resume.docx", resume, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers={"Authorization": f"Bearer {candidate_token}"},
        )
    assert response.status_code == 400
    assert "already applied" in response.json()["detail"]


def test_invalid_resume(hr_token, candidate_token):
    job_response = client.post(
        "/jobs",
        json={"title": "Another Role", "description": "desc", "requirements": "reqs"},
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    job_id = job_response.json()["id"]
    response = client.post(
        f"/jobs/{job_id}/apply",
        files={"resume": ("notes.txt", b"plain text, not a resume", "text/plain")},
        headers={"Authorization": f"Bearer {candidate_token}"},
    )
    assert response.status_code == 400
    assert "Only PDF and DOCX" in response.json()["detail"]


def test_resume_parsing(sample_resume):
    text = extract_text_from_file(str(sample_resume))
    assert "Alex Doe" in text
    assert "Python" in text
    assert "FastAPI" in text


@pytest.mark.skipif(not os.getenv("GROQ_API_KEY"), reason="GROQ_API_KEY not set")
def test_ai_scoring_response():
    from ai_scorer import score_resume_against_job
    resume = "Python developer with 3 years of FastAPI, SQL, REST API experience."
    description = "Backend developer role building APIs."
    requirements = "Python, FastAPI, SQL, REST API design"
    result = score_resume_against_job(resume, description, requirements)
    assert isinstance(result.get("match_score"), (int, float))
    assert 0 <= result["match_score"] <= 100
    assert isinstance(result.get("matched_skills"), list)
    assert isinstance(result.get("missing_skills"), list)
    assert result.get("explanation")


@pytest.mark.skipif(not os.getenv("GROQ_API_KEY"), reason="GROQ_API_KEY not set")
def test_rag_retrieval():
    from rag_assistant import ask_candidate_rag
    result = ask_candidate_rag(
        resume_text="Alex Doe. Python developer. Built RAG pipelines with LangChain.",
        candidate_name="Alex Doe",
        candidate_meta={"match_score": 80, "skills": "['Python']", "experience": "3 years", "education": "B.E."},
        job_title="Backend Engineer",
        job_description="Building APIs",
        job_requirements="Python, LLMs",
        question="Does the candidate have RAG experience?",
    )
    assert isinstance(result.get("answer"), str)
    assert result["answer"]
    assert "suggested_followups" in result


def test_unauthorized_endpoint():
    response = client.get("/jobs/mine")
    assert response.status_code == 401


def test_candidate_cannot_create_job(candidate_token):
    response = client.post(
        "/jobs",
        json={"title": "Nope", "description": "desc", "requirements": "reqs"},
        headers={"Authorization": f"Bearer {candidate_token}"},
    )
    assert response.status_code == 403


def test_interview_slot_creation(job_id, hr_token, candidate_token, sample_resume):
    with open(sample_resume, "rb") as resume:
        application = client.post(
            f"/jobs/{job_id}/apply",
            files={"resume": ("resume.docx", resume, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers={"Authorization": f"Bearer {candidate_token}"},
        ).json()
    application_id = application["id"]

    shortlist = client.put(
        f"/applications/{application_id}/status",
        json={"status": "shortlisted"},
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert shortlist.status_code == 200

    response = client.post(
        f"/applications/{application_id}/interview-slots",
        json={"scheduled_at": ["2026-10-01T10:00:00", "2026-10-02T11:00:00"]},
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    assert response.status_code == 200
    slots = response.json()
    assert len(slots) == 2
    assert slots[0]["status"] == "proposed"


def test_candidate_slot_selection(job_id, hr_token, candidate_token, sample_resume):
    with open(sample_resume, "rb") as resume:
        application = client.post(
            f"/jobs/{job_id}/apply",
            files={"resume": ("resume.docx", resume, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers={"Authorization": f"Bearer {candidate_token}"},
        ).json()
    application_id = application["id"]
    client.put(
        f"/applications/{application_id}/status",
        json={"status": "shortlisted"},
        headers={"Authorization": f"Bearer {hr_token}"},
    )
    slot_id = client.post(
        f"/applications/{application_id}/interview-slots",
        json={"scheduled_at": ["2026-10-01T10:00:00"]},
        headers={"Authorization": f"Bearer {hr_token}"},
    ).json()[0]["id"]

    response = client.post(
        f"/interview-slots/{slot_id}/select",
        headers={"Authorization": f"Bearer {candidate_token}"},
    )
    assert response.status_code == 200
    assert response.json()["status"] == "scheduled"